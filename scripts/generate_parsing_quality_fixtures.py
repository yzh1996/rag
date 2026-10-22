"""生成小而可复现的文档解析测试资料，不调用 OCR 或大模型。

每份文件旁都有 golden.json（人工设定的正确答案）。答案来自生成前的
样例设计，绝不从解析器的输出反推，否则解析错了也可能被当成正确。
"""
from __future__ import annotations

import argparse
import json
import os
from pathlib import Path
import shutil
import subprocess

import fitz
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'tests' / '测试资源' / '文档解析质量'
FONT = Path(os.environ.get('PARSING_FIXTURE_FONT', r'C:\Windows\Fonts\msyh.ttc'))
FIELDS = ['订单编号：PO-2026-000108', '核销金额：-1000.05', '日期：2026-10-21',
          '姓名：欧阳明', '字母标记：OOO', '处理结论：不予报销']
CASES: list[dict] = []


def check(kind, name, **kwargs):
    return {'kind': kind, 'name': name, **kwargs}


def required(lines, *, ocr=False):
    return [check('contains', line.split('：')[0], value=line, requires_ocr=ocr) for line in lines]


def case(identifier, category, filename, description, checks, **kwargs):
    directory = OUT / category / identifier
    directory.mkdir(parents=True, exist_ok=True)
    record = dict(id=identifier, category=category, file=filename, description=description,
                  loader='pymupdf', checks=checks, **kwargs)
    CASES.append(record)
    return directory / filename, record


def pdf(path, pages):
    """pages 是绘图函数列表；顺序就是正确的物理页顺序。"""
    canvas = Canvas(str(path), pagesize=A4, invariant=1)
    canvas.setTitle(path.stem)
    for draw in pages:
        canvas.setFillColorRGB(0.08, 0.13, 0.22)
        canvas.setFont('FixtureChinese', 12)
        draw(canvas)
        canvas.showPage()
    canvas.save()


def lines(canvas, values, x=42, top=65, step=34):
    for index, value in enumerate(values):
        canvas.drawString(x, A4[1] - top - index * step, value)


def scan_image(mode='clear'):
    image = Image.new('RGB', (1240, 1754), 'white')
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(str(FONT), 32)
    for index, value in enumerate(['扫描件关键字段测试'] + FIELDS):
        draw.text((80, 110 + index * 80), value, font=font, fill='#182235')
    if mode == 'low_resolution':
        image = image.resize((372, 526), Image.Resampling.BILINEAR)
    elif mode == 'tilted':
        image = image.rotate(4, resample=Image.Resampling.BICUBIC, expand=False, fillcolor='white')
    elif mode == 'stamp':
        draw.ellipse((560, 300, 870, 600), outline='#CC3333', width=8)
        draw.text((620, 410), '测试专用', font=font, fill='#CC3333')
    return image


def image_page(canvas, image):
    canvas.drawImage(ImageReader(image), 0, 0, width=A4[0], height=A4[1])


def table(canvas, rows, *, top=110, borders=True, merge=False):
    xs = [42, 170, 310, 460]
    row_height = 40
    y = A4[1] - top
    if borders:
        for x in xs:
            canvas.line(x, y, x, y - row_height * len(rows))
        for index in range(len(rows) + 1):
            # 合并 A2:A3：中间横线从第二列才开始，不能用两个空单元格假装合并。
            left = xs[1] if merge and index == 2 else xs[0]
            canvas.line(left, y - row_height * index, xs[-1], y - row_height * index)
    for row_index, row in enumerate(rows):
        for col, value in enumerate(row):
            canvas.drawString(xs[col] + 8, y - row_index * row_height - 25, str(value))


def create_accuracy():
    path, _ = case('text_fields', '01_文字准确性', '数字与易混字符.pdf',
                   '真实文字层：0/O、负号、小数点、日期、姓名和否定词。', required(FIELDS) + [
                       check('cer', '关键字段字符错误率', expected='\n'.join(FIELDS), max_rate=0.0)])
    pdf(path, [lambda c: lines(c, FIELDS)])
    for mode, title in [('clear', '清晰'), ('low_resolution', '低分辨率'), ('tilted', '倾斜'), ('stamp', '模拟盖章')]:
        path, _ = case(f'scan_{mode}', '01_文字准确性', f'{title}扫描件.pdf',
                       f'{title}图片页，无隐藏文字层；需真实 OCR 后才能评价字段准确率。',
                       [check('ocr_detection', '扫描页识别', pages=[1])] + required(FIELDS, ocr=True), ocr_pages=[1])
        image = scan_image(mode)
        pdf(path, [lambda c, image=image: image_page(c, image)])
    path, _ = case('mixed_pages', '02_完整性', '文字页与扫描页混合.pdf',
                   '第一、三页有文字，第二页只有图片；数字页也要保留。',
                   required(['数字页开始：保留采购政策的适用范围和责任部门。', '数字页结束：提交财务审核。']) +
                   [check('ocr_detection', '仅第二页需要 OCR', pages=[2])] + required(FIELDS, ocr=True), ocr_pages=[2])
    pdf(path, [lambda c: lines(c, ['数字页开始：保留采购政策的适用范围和责任部门。']),
               lambda c: image_page(c, scan_image()), lambda c: lines(c, ['数字页结束：提交财务审核。'])])
    path, _ = case('mixed_same_page', '02_完整性', '同一页文字加图片.pdf',
                   '已有足够长的正文，但关键金额只在图片里；检查是否漏掉局部图片。',
                   required(['本页正文已超过扫描检测阈值，仍需要识别下面图片中的金额和编号。']) +
                   [check('ocr_detection', '同页图片不能漏检', pages=[1])] + required(FIELDS, ocr=True), ocr_pages=[1])
    def mixed(c):
        lines(c, ['本页正文已超过扫描检测阈值，仍需要识别下面图片中的金额和编号。'], step=25)
        c.drawImage(ImageReader(scan_image()), 60, 70, width=420, height=594)
    pdf(path, [mixed])


def create_order():
    markers = ['LEFT-A', 'LEFT-B', 'RIGHT-A', 'RIGHT-B']
    path, _ = case('columns', '03_阅读顺序', '双栏交错存储.pdf',
                   '人眼应先读完左栏再读右栏；PDF 内部故意按横向交错写入。',
                   [check('order', '先左栏再右栏', values=markers)])
    def columns(c):
        lines(c, ['双栏阅读测试：先左后右'])
        for text, x, top in [('LEFT-A 左栏第一段', 42, 140), ('RIGHT-A 右栏第一段', 320, 140),
                             ('LEFT-B 左栏第二段', 42, 220), ('RIGHT-B 右栏第二段', 320, 220)]:
            lines(c, [text], x=x, top=top)
    pdf(path, [columns])
    path, _ = case('sidebar_footnote', '03_阅读顺序', '侧栏脚注.pdf',
                   '正文先连续阅读，补充说明和脚注放在正文之后。',
                   [check('order', '正文不被侧栏打断', values=['MAIN-A', 'MAIN-B', 'SIDE', 'FOOT'])])
    def sidebar(c):
        lines(c, ['MAIN-A 正文第一段'], top=100)
        lines(c, ['SIDE 补充说明'], x=365, top=120)
        lines(c, ['MAIN-B 正文第二段'], top=180)
        lines(c, ['FOOT 脚注：金额单位为元。'], top=740)
    pdf(path, [sidebar])
    path, _ = case('cross_page_text', '03_阅读顺序', '跨页段落.pdf',
                   '段落在页尾中断，下一页继续；页码需能追溯。',
                   [check('order', '跨页连续', values=['START', 'CONTINUE', 'END']),
                    check('pages', '来源页码', values=[1, 2])])
    pdf(path, [lambda c: lines(c, ['START 本条款在下一页继续。'], top=740),
               lambda c: lines(c, ['CONTINUE 审批完成后归档。', 'END 条款结束。'])])


def create_tables():
    basic = [['商品', '销量', '金额'], ['商品甲', '10', '1000.05'], ['商品乙', '2', '-20.50']]
    for identifier, name, borders in [('table_bordered', '有边框表格', True), ('table_borderless', '无边框表格', False)]:
        path, record = case(identifier, '04_表格关系', name + '.pdf',
                            '必须保留表头、行和数值的对应关系，不能只判断所有词是否出现。',
                            [check('table', '商品、销量、金额对应', rows=basic)])
        record['loader'] = 'pdfplumber'
        pdf(path, [lambda c, borders=borders: table(c, basic, borders=borders)])
    merged = [['部门', '商品', '金额'], ['华东', '商品甲', '1000.05'], ['', '商品乙', '-20.50']]
    path, record = case('table_merged', '04_表格关系', '合并单元格.pdf',
                        '华东是商品甲和商品乙共同的部门；提取空白单元格不能丢失归属。',
                        [check('table', '合并单元格归属', rows=[merged[0], merged[1], ['华东', '商品乙', '-20.50']])])
    record['loader'] = 'pdfplumber'
    pdf(path, [lambda c: table(c, merged, merge=True)])
    path, record = case('table_cross_page', '04_表格关系', '跨页重复表头.pdf',
                        '两页是同一张表，需要保留两条记录以及各自页码。',
                        [check('table', '跨页记录对应', rows=basic), check('pages', '表格来源页', values=[1, 2])])
    record['loader'] = 'pdfplumber'
    pdf(path, [lambda c: table(c, basic[:2]), lambda c: table(c, [basic[0], basic[2]])])
    for complex in [False, True]:
        name = '多表合并空行.xlsx' if complex else '基础表格.xlsx'
        checks = required(['000108', '000109']) + [check('table', 'Excel 行列对应', rows=[
            ['部门', '编号', '商品', '金额'], ['华东', '000108', '商品甲', '1000.05'],
            ['华东', '000109', '商品乙', '-20.5']
        ] if complex else [['编号', '商品', '数量', '金额'], ['000108', '商品甲', '10', '1000.05'],
                           ['000109', '商品乙', '2', '-20.5']])]
        if complex:
            checks.append(check('metadata_values', '工作表来源', key='page_name', values=['一月', '二月']))
        case('excel_complex' if complex else 'excel_basic', '04_表格关系', name,
             '保留编号前导零；复杂版另含两张工作表、合并单元格和空白行。', checks)


def create_structure_cleaning():
    policy = ['采购管理办法', '版本：V2.1', '生效日期：2026-10-01', '责任部门：采购部', '第一章 适用范围', '本办法适用于正式采购申请。']
    path, _ = case('policy_metadata', '05_结构来源', '制度版本与来源.pdf',
                   '版本和生效日期取自正文；部门、密级由上传人提供，不能当成解析器自动识别。',
                   required(policy) + [check('pages', '来源页', values=[1]),
                    check('source', '来源文件', filename='制度版本与来源.pdf'),
                    check('metadata_values', '用户提供的部门', key='department', values=['采购部'], provided_by_user=True)],
                   input_metadata={'department': '采购部', 'security_level': '内部', 'knowledge_base': '采购制度'})
    pdf(path, [lambda c: lines(c, policy)])
    path, _ = case('docx_headings', '05_结构来源', '标题层级.docx',
                   '不仅要保留标题文字，还要能区分一级与二级标题。',
                   required(['费用管理手册', '一、报销', '1.1 差旅']) +
                   [check('heading', '一级标题层级', text='一、报销', level=1),
                    check('heading', '二级标题层级', text='1.1 差旅', level=2)])
    doc = word_document()
    doc.add_heading('费用管理手册', level=0)
    doc.add_heading('一、报销', level=1)
    doc.add_heading('1.1 差旅', level=2)
    doc.add_paragraph('应保留发票与审批编号。')
    doc.save(path)
    path, _ = case('docx_noise', '06_清洗与异常', '页眉空行免责声明.docx',
                   '只删除本样例已知噪声，同时保留编号、金额、否定词。', required(FIELDS) +
                   [check('absent', '删除页眉', value='测试企业内部页眉'),
                    check('absent', '删除样例免责声明', value='免责声明：本文件仅供测试。'),
                    check('normalized', '连续空白规范化')])
    doc = word_document()
    doc.sections[0].header.paragraphs[0].text = '测试企业内部页眉'
    doc.add_paragraph('免责声明：本文件仅供测试。')
    for value in FIELDS:
        doc.add_paragraph(value)
    doc.add_paragraph('审批    责任人\t财务部')
    for _ in range(4):
        doc.add_paragraph('')
    doc.add_paragraph('结束。')
    doc.save(path)
    path, _ = case('repeated_edges', '06_清洗与异常', '重复页眉页脚.pdf',
                   '三页相同的页眉页脚；报告会记录当前是否启用了重复边缘清洗。',
                   required(['正文第一部分', '正文第二部分', '正文第三部分']) +
                   [check('absent', '删除重复页眉', value='测试企业机密页眉'), check('absent', '删除重复页脚', value='测试企业机密页脚')])
    def page(c, text):
        lines(c, ['测试企业机密页眉'], top=35)
        lines(c, [text], top=150)
        lines(c, ['测试企业机密页脚'], top=790)
    pdf(path, [lambda c, text=text: page(c, text) for text in ['正文第一部分', '正文第二部分', '正文第三部分']])
    path, _ = case('blank_page', '06_清洗与异常', '中间空白页.pdf',
                   '允许丢弃空白内容，但第三页不能被误标成第二页。', required(['第一页有正文。', '第三页有正文。']) +
                   [check('nonempty', '清洗后没有空内容单元'), check('pages', '空白页后仍保留真实页码', values=[1, 3])])
    pdf(path, [lambda c: lines(c, ['第一页有正文。']), lambda c: None, lambda c: lines(c, ['第三页有正文。'])])


def word_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.8)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
        style = doc.styles[name]
        style.font.name = 'Microsoft YaHei'
        style.font.color.rgb = RGBColor(0, 0, 0)
        if name == 'Normal':
            style.font.size = Pt(11)
    return doc


def create_rejections():
    path, record = case('corrupt_pdf', '06_清洗与异常', '损坏文件.pdf',
                        '伪装成 PDF 的损坏文件，应该拒绝且不能产出可用内容。', [])
    path.write_bytes(b'%PDF-1.7\nintentionally incomplete test fixture\n')
    record.update(expected_rejection=True, rejection_types=['FileDataError', 'FzErrorFormat', 'PDFSyntaxError', 'ValueError'])
    path, record = case('encrypted_pdf', '06_清洗与异常', '加密文件.pdf', '已加密文件应提示先解密。', [])
    with fitz.open() as document:
        document.new_page().insert_text((30, 40), 'Encrypted fixture')
        document.save(path, encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw='fixture-owner', user_pw='fixture-reader')
    record.update(expected_rejection=True, error_contains='加密PDF', rejection_types=['ValueError', 'PDFPasswordIncorrect'])
    path, record = case('page_limit', '06_清洗与异常', '四页超限样例.pdf',
                        '测试时暂把页数限制设为 3，用 4 个小页面验证拒绝，不创建大文件。', [])
    pdf(path, [lambda c, i=i: lines(c, [f'LIMIT-{i} 小文件资源限制测试']) for i in range(4)])
    record.update(expected_rejection=True, error_contains='页上限', rejection_types=['ValueError'], settings={'MAX_PDF_PAGES': 3})
    path, record = case('office_limit', '06_清洗与异常', 'Office解压限制.docx',
                        '普通小型 Word；测试时限制 ZIP 条目数为 1，验证解压前保护。', [])
    doc = word_document()
    doc.add_paragraph('普通的小型 Word 文件，不包含危险压缩数据。')
    doc.save(path)
    record.update(expected_rejection=True, error_contains='解压后过大', rejection_types=['ValueError'], settings={'MAX_ARCHIVE_ENTRIES': 1})


def create_excel(node):
    stage = OUT / '_excel_preview'
    subprocess.run([node, str(Path(__file__).with_name('generate_parsing_quality_excel.mjs')), str(stage)], check=True)
    for record in CASES:
        if record['file'].endswith('.xlsx'):
            shutil.copyfile(stage / record['file'], OUT / record['category'] / record['id'] / record['file'])


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--node', default='node', help='安装 @oai/artifact-tool 的 Node 可执行文件')
    args = parser.parse_args()
    if not FONT.exists():
        raise SystemExit('未找到中文字体，可用 PARSING_FIXTURE_FONT 指定字体路径。')
    pdfmetrics.registerFont(TTFont('FixtureChinese', str(FONT)))
    create_accuracy()
    create_order()
    create_tables()
    create_structure_cleaning()
    create_rejections()
    create_excel(args.node)
    for record in CASES:
        (OUT / record['category'] / record['id'] / 'golden.json').write_text(
            json.dumps(record, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    (OUT / 'manifest.json').write_text(json.dumps({
        'version': 1, 'note': '全部为虚构测试数据；golden.json 是独立设定的标准答案。',
        'cases': [f"{c['category']}/{c['id']}/golden.json" for c in CASES],
    }, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f'已生成 {len(CASES)} 个分类样例：{OUT}')


if __name__ == '__main__':
    main()
