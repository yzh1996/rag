"""从原始测试文件生成报告预览，供人眼对照，不作为解析结果或评分答案。"""
from __future__ import annotations

from html import escape
from pathlib import Path
import shutil
import xml.etree.ElementTree as ET
from zipfile import ZipFile

WORD_NS = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def _word_text(paragraph):
    parts = []
    for node in paragraph.iter():
        if node.tag == WORD_NS + 't':
            parts.append(node.text or '')
        elif node.tag == WORD_NS + 'tab':
            parts.append('\t')
        elif node.tag in (WORD_NS + 'br', WORD_NS + 'cr'):
            parts.append('\n')
    return ''.join(parts)


def _word_preview(path):
    # 直接读取原件的段落/样式；保留页眉和空白，不能用“清洗后文字”冒充原件。
    chunks = ['<p class="preview-note">Word 结构预览，保留标题样式；不代表 Word 的原始分页。</p>']
    with ZipFile(path) as archive:
        for name in sorted(archive.namelist()):
            if name.startswith('word/header') and name.endswith('.xml'):
                header = ET.fromstring(archive.read(name))
                chunks.append('<div class="word-edge">页眉：' + escape(_word_text(header)) + '</div>')
        body = ET.fromstring(archive.read('word/document.xml')).find(WORD_NS + 'body')
        for paragraph in body.findall(WORD_NS + 'p'):
            style = paragraph.find(f'{WORD_NS}pPr/{WORD_NS}pStyle')
            style_name = style.get(WORD_NS + 'val', '') if style is not None else ''
            css = {'Title': 'word-title', 'Heading1': 'word-heading', 'Heading2': 'word-subheading'}.get(style_name, 'word-paragraph')
            chunks.append(f'<div class="{css}">{escape(_word_text(paragraph)) or "&nbsp;"}</div>')
        for name in sorted(archive.namelist()):
            if name.startswith('word/footer') and name.endswith('.xml'):
                chunks.append('<div class="word-edge">页脚：' + escape(_word_text(ET.fromstring(archive.read(name)))) + '</div>')
    return ''.join(chunks)


def _excel_preview(path):
    from openpyxl import load_workbook

    chunks = ['<p class="preview-note">Excel 单元格结构预览，保留工作表、合并范围和空白行。</p>']
    book = load_workbook(path, data_only=False)
    try:
        for sheet in book:
            chunks.append(f'<h4>工作表：{escape(sheet.title)}</h4><div class="sheet-scroll"><table class="source-table">')
            spans, covered = {}, set()
            for merged in sheet.merged_cells.ranges:
                spans[(merged.min_row, merged.min_col)] = (merged.max_row - merged.min_row + 1, merged.max_col - merged.min_col + 1)
                covered.update((row, col) for row in range(merged.min_row, merged.max_row + 1)
                               for col in range(merged.min_col, merged.max_col + 1)
                               if (row, col) != (merged.min_row, merged.min_col))
            # 当前样例都很小；将来换成大表时，避免报告预览占用过多资源。
            for row in sheet.iter_rows(max_row=min(sheet.max_row, 200), max_col=min(sheet.max_column, 30)):
                chunks.append('<tr>')
                for cell in row:
                    if (cell.row, cell.column) in covered:
                        continue
                    height, width = spans.get((cell.row, cell.column), (1, 1))
                    value = '' if cell.value is None else str(cell.value)
                    if isinstance(cell.value, (int, float)) and cell.number_format == '0.00':
                        value = f'{cell.value:.2f}'
                    chunks.append(f'<td rowspan="{height}" colspan="{width}">{escape(value) or "&nbsp;"}</td>')
                chunks.append('</tr>')
            chunks.append('</table></div>')
            if sheet.max_row > 200 or sheet.max_column > 30:
                chunks.append('<p>预览仅展示前 200 行、30 列，请打开原件查看完整内容。</p>')
    finally:
        book.close()
    return ''.join(chunks)


def build_original_preview(case, target):
    """预览失败只影响展示，不阻止两种解析结果和失败报告保存。"""
    source = Path(case['path'])
    result = {'status': 'unavailable', 'html': '', 'pages': [], 'source_href': '', 'golden_href': ''}
    try:
        target.mkdir(parents=True, exist_ok=True)
        local_source = target / ('source' + source.suffix.lower())
        shutil.copyfile(source, local_source)
        result['source_href'] = f'{target.name}/{local_source.name}'
        golden = Path(case['golden_path'])
        if golden.exists():
            shutil.copyfile(golden, target / 'golden.json')
            result['golden_href'] = f'{target.name}/golden.json'
        if source.suffix.lower() == '.pdf':
            import fitz

            with fitz.open(source) as pdf:
                if pdf.is_encrypted:
                    result['html'] = '<p>原文件已加密，无法直接预览；请打开原件并解密后查看。</p>'
                    return result
                for number, page in enumerate(pdf, 1):
                    name = f'original-page-{number:03d}.png'
                    page.get_pixmap(dpi=110, alpha=False).save(target / name)
                    result['pages'].append({'page': number, 'image': f'{target.name}/{name}'})
            result['html'] = ''.join(
                f'<figure><figcaption>原件第 {page["page"]} 页</figcaption>'
                f'<a href="{escape(page["image"])}" target="_blank"><img loading="lazy" '
                f'src="{escape(page["image"])}" alt="原件第 {page["page"]} 页"></a></figure>'
                for page in result['pages'])
        elif source.suffix.lower() == '.docx':
            result['html'] = _word_preview(source)
        elif source.suffix.lower() == '.xlsx':
            result['html'] = _excel_preview(source)
        else:
            result['html'] = '<p>此格式请通过原始文件链接查看。</p>'
            return result
        result['status'] = 'ready'
    except Exception as exc:
        result['html'] = f'<p>原件预览不可用（{escape(type(exc).__name__)}），不影响解析检查结果。</p>'
    return result


def render_original_panel(case, original):
    links = []
    if original['source_href']:
        links.append(f'<a href="{escape(original["source_href"])}" target="_blank">打开原始文件</a>')
    if original['golden_href']:
        links.append(f'<a href="{escape(original["golden_href"])}" target="_blank">标准答案 JSON</a>')
    chunks = ['<section class="route original"><h3>原始资料与标准答案</h3>',
              '<p>直接查看原件和预先设定的答案，作为人工核对依据。</p>', '<p>' + ' · '.join(links) + '</p>',
              '<details class="golden"><summary>查看标准答案与检查要求</summary><table><tr><th>检查项</th><th>标准</th></tr>']
    for rule in case['checks']:
        value = next((rule[key] for key in ['value', 'expected', 'rows', 'values', 'pages', 'filename', 'text'] if key in rule), rule['name'])
        if rule['kind'] == 'table':
            displayed = '<table class="source-table">' + ''.join('<tr>' + ''.join(f'<td>{escape(str(cell))}</td>' for cell in row) + '</tr>' for row in value) + '</table>'
        elif isinstance(value, list):
            displayed = escape(' → '.join(str(item) for item in value))
        else:
            displayed = escape(str(value))
        chunks.append(f'<tr><td>{escape(rule["name"])}</td><td>{displayed}</td></tr>')
    if case.get('expected_rejection'):
        chunks.append('<tr><td>预期处理</td><td>拒绝该文件：' + escape(case.get('error_contains') or '文件损坏') + '</td></tr>')
    chunks.extend(['</table></details><div class="source-preview">', original['html'], '</div></section>'])
    return ''.join(chunks)
